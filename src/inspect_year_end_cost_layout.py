from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "year_end_inventory"
    / "dec 2022.xlsx"
)


def merged_range_for_cell(
    worksheet,
    row_number,
    column_number,
):
    for merged_range in worksheet.merged_cells.ranges:
        if (
            merged_range.min_row
            <= row_number
            <= merged_range.max_row
            and
            merged_range.min_col
            <= column_number
            <= merged_range.max_col
        ):
            return str(merged_range)

    return ""


def main():
    print("=" * 90)
    print("CDH INVENTORY DSS")
    print("YEAR-END COST FIELD DIAGNOSTIC")
    print("=" * 90)

    # One workbook gives us formulas.
    workbook_formula = load_workbook(
        FILE_PATH,
        read_only=False,
        data_only=False,
    )

    # The second gives us cached/displayed values.
    workbook_value = load_workbook(
        FILE_PATH,
        read_only=False,
        data_only=True,
    )

    worksheet_formula = workbook_formula[
        workbook_formula.sheetnames[0]
    ]

    worksheet_value = workbook_value[
        workbook_value.sheetnames[0]
    ]

    print()
    print("File:", FILE_PATH.name)
    print("Sheet:", worksheet_formula.title)

    # ========================================================
    # PART 1 — COST-AREA HEADER
    # ========================================================

    print()
    print("-" * 90)
    print("COST / BALANCE AREA — ROWS 20 TO 38")
    print("COLUMNS Y TO AO")
    print("-" * 90)

    for row_number in range(20, 39):

        found = []

        for column_number in range(25, 42):
            formula_value = worksheet_formula.cell(
                row=row_number,
                column=column_number,
            ).value

            cached_value = worksheet_value.cell(
                row=row_number,
                column=column_number,
            ).value

            if (
                formula_value is None
                and cached_value is None
            ):
                continue

            coordinate = (
                f"{get_column_letter(column_number)}"
                f"{row_number}"
            )

            merged_range = (
                merged_range_for_cell(
                    worksheet_formula,
                    row_number,
                    column_number,
                )
            )

            found.append(
                {
                    "coordinate": coordinate,
                    "formula": formula_value,
                    "cached": cached_value,
                    "merged": merged_range,
                }
            )

        if found:
            print()
            print(f"ROW {row_number}")

            for item in found:
                print(
                    f"  {item['coordinate']}"
                )
                print(
                    f"    formula/raw : "
                    f"{item['formula']}"
                )
                print(
                    f"    cached value: "
                    f"{item['cached']}"
                )

                if item["merged"]:
                    print(
                        f"    merged range: "
                        f"{item['merged']}"
                    )

    # ========================================================
    # PART 2 — FIRST FIVE MEDICINE ROWS
    # ========================================================

    print()
    print("-" * 90)
    print("FIRST FIVE MEDICINE ROWS — ALL NONEMPTY CELLS")
    print("-" * 90)

    for row_number in range(34, 39):

        print()
        print(f"ROW {row_number}")

        for column_number in range(
            1,
            worksheet_formula.max_column + 1,
        ):

            formula_value = worksheet_formula.cell(
                row=row_number,
                column=column_number,
            ).value

            cached_value = worksheet_value.cell(
                row=row_number,
                column=column_number,
            ).value

            if (
                formula_value is None
                and cached_value is None
            ):
                continue

            coordinate = (
                f"{get_column_letter(column_number)}"
                f"{row_number}"
            )

            merged_range = (
                merged_range_for_cell(
                    worksheet_formula,
                    row_number,
                    column_number,
                )
            )

            print(
                f"  {coordinate}"
            )

            print(
                f"    formula/raw : "
                f"{formula_value}"
            )

            print(
                f"    cached value: "
                f"{cached_value}"
            )

            if merged_range:
                print(
                    f"    merged range: "
                    f"{merged_range}"
                )

    print()
    print("=" * 90)
    print("DIAGNOSTIC COMPLETE")
    print("=" * 90)
    print(
        "No workbook value was modified."
    )


if __name__ == "__main__":
    main()