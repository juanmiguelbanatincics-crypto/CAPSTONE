from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# ============================================================
# PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

YEAR_END_DIR = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "year_end_inventory"
)

ANNUAL_DIR = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "annual_inventory"
)

OUTPUT_DIR = (
    PROJECT_ROOT
    / "data"
    / "processed"
    / "year_end_reconciliation"
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# ============================================================
# ACTUAL PHYSICAL DATA COLUMNS
# CONFIRMED FROM CDH BIZBOX EXCEL EXPORT
# ============================================================

# Excel column:
#
# B  = Item ID
# K  = Generic Name
# R  = Item Description
# W  = Unit
# AA = Last Purchase Price
# AI = Balance
# AL = Total Inventory Value
#
# The report contains a "Unit Cost" header,
# but the corresponding medicine-row cells
# are not populated in the source workbook.

COLUMN_MAP = {
    "Item_ID": "B",
    "Generic_Name_YE": "K",
    "Item_Description": "R",
    "Unit_YE": "W",
    "Last_Purchase_Price": "AA",
    "Balance_YE": "AI",
    "Total_Value_YE": "AL",
}


# ============================================================
# HELPERS
# ============================================================

def infer_year(
    filename: str,
) -> int:

    match = re.search(
        r"(20\d{2})",
        filename,
    )

    if not match:
        raise ValueError(
            f"Could not infer year from "
            f"filename: {filename}"
        )

    return int(
        match.group(1)
    )


def clean_text(
    value,
):

    if value is None:
        return pd.NA

    text = str(value).strip()

    if not text:
        return pd.NA

    return text


def normalize_match_text(
    series: pd.Series,
) -> pd.Series:

    return (
        series
        .astype("string")
        .str.upper()
        .str.strip()
        .str.replace(
            r"\s+",
            " ",
            regex=True,
        )
        .str.replace(
            r"[^A-Z0-9]+",
            "",
            regex=True,
        )
        .replace(
            "",
            pd.NA,
        )
    )


def percent(
    numerator: int | float,
    denominator: int | float,
) -> float | None:

    if denominator == 0:
        return None

    return round(
        numerator
        / denominator
        * 100,
        4,
    )


def read_annual_csv(
    path: Path,
) -> pd.DataFrame:

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            return pd.read_csv(
                path,
                encoding=encoding,
                low_memory=False,
            )

        except UnicodeDecodeError as error:
            last_error = error

    raise last_error


# ============================================================
# PARSE YEAR-END WORKBOOK
# ============================================================

def parse_year_end_workbook(
    path: Path,
) -> pd.DataFrame:

    year = infer_year(
        path.name
    )

    workbook = load_workbook(
        path,
        read_only=False,
        data_only=True,
    )

    worksheet = workbook[
        workbook.sheetnames[0]
    ]

    records = []

    # Physical medicine records begin
    # after the formatted report header.
    #
    # We do not depend on one exact first
    # row. Instead, rows are accepted only
    # when an Item ID beginning with MED
    # is present.

    for row_number in range(
        28,
        worksheet.max_row + 1,
    ):

        item_id = clean_text(
            worksheet[
                f"{COLUMN_MAP['Item_ID']}"
                f"{row_number}"
            ].value
        )

        if pd.isna(item_id):
            continue

        item_id_string = str(
            item_id
        ).strip()

        if not (
            item_id_string
            .upper()
            .startswith("MED")
        ):
            continue

        record = {
            "Year": year,
            "Source_File": path.name,
            "Excel_Row": row_number,
        }

        for (
            field,
            column_letter,
        ) in COLUMN_MAP.items():

            value = worksheet[
                f"{column_letter}"
                f"{row_number}"
            ].value

            record[field] = value

        records.append(
            record
        )

    dataframe = pd.DataFrame(
        records
    )

    if dataframe.empty:
        raise ValueError(
            f"No MED item rows parsed "
            f"from {path.name}"
        )

    # --------------------------------------------------------
    # CLEAN TEXT
    # --------------------------------------------------------

    for field in [
        "Item_ID",
        "Generic_Name_YE",
        "Item_Description",
        "Unit_YE",
    ]:

        dataframe[field] = (
            dataframe[field]
            .map(clean_text)
            .astype("string")
        )

    # --------------------------------------------------------
    # NUMERIC FIELDS
    # --------------------------------------------------------

    for field in [
        "Last_Purchase_Price",
        "Balance_YE",
        "Total_Value_YE",
    ]:

        dataframe[field] = (
            pd.to_numeric(
                dataframe[field],
                errors="coerce",
            )
        )

    # --------------------------------------------------------
    # DERIVED IMPLIED VALUATION COST
    # --------------------------------------------------------

    valid_implied_cost = (
        dataframe[
            "Balance_YE"
        ].notna()
        &
        dataframe[
            "Total_Value_YE"
        ].notna()
        &
        dataframe[
            "Balance_YE"
        ].ne(0)
    )

    dataframe[
        "Implied_Valuation_Cost"
    ] = np.nan

    dataframe.loc[
        valid_implied_cost,
        "Implied_Valuation_Cost",
    ] = (
        dataframe.loc[
            valid_implied_cost,
            "Total_Value_YE",
        ]
        /
        dataframe.loc[
            valid_implied_cost,
            "Balance_YE",
        ]
    )

    # --------------------------------------------------------
    # NORMALIZED MATCHING FIELDS
    # --------------------------------------------------------

    dataframe[
        "Normalized_Description"
    ] = normalize_match_text(
        dataframe[
            "Item_Description"
        ]
    )

    dataframe[
        "Normalized_Generic_YE"
    ] = normalize_match_text(
        dataframe[
            "Generic_Name_YE"
        ]
    )

    dataframe[
        "Normalized_Unit_YE"
    ] = normalize_match_text(
        dataframe[
            "Unit_YE"
        ]
    )

    return dataframe


# ============================================================
# LOAD ANNUAL MEDICINES
# ============================================================

def load_annual_medicines(
    year: int,
) -> pd.DataFrame:

    candidates = list(
        ANNUAL_DIR.glob(
            f"*{year}*.csv"
        )
    )

    if len(candidates) != 1:
        raise ValueError(
            f"Expected one annual CSV "
            f"for {year}; found "
            f"{len(candidates)}"
        )

    dataframe = read_annual_csv(
        candidates[0]
    )

    dataframe.columns = [
        str(column).strip()
        for column
        in dataframe.columns
    ]

    medicines = (
        dataframe.loc[
            dataframe[
                "Category1"
            ]
            .astype("string")
            .str.strip()
            .str.upper()
            .eq("MEDICINES")
        ]
        .copy()
    )

    for field in [
        "Brand_Name",
        "Generic_Name",
        "Unit",
    ]:

        medicines[field] = (
            medicines[field]
            .astype("string")
            .str.strip()
        )

    for field in [
        "Ending_Balance",
        "Average_cost",
        "Qty_out_Stock_Issuances",
    ]:

        medicines[field] = (
            pd.to_numeric(
                medicines[field],
                errors="coerce",
            )
        )

    medicines[
        "Normalized_Brand"
    ] = normalize_match_text(
        medicines[
            "Brand_Name"
        ]
    )

    medicines[
        "Normalized_Generic_Annual"
    ] = normalize_match_text(
        medicines[
            "Generic_Name"
        ]
    )

    medicines[
        "Normalized_Unit_Annual"
    ] = normalize_match_text(
        medicines[
            "Unit"
        ]
    )

    medicines["Year"] = year

    return medicines


# ============================================================
# YEAR-END PROFILE
# ============================================================

def profile_year_end(
    dataframe: pd.DataFrame,
) -> dict:

    return {
        "year":
            int(
                dataframe[
                    "Year"
                ].iloc[0]
            ),

        "source_file":
            dataframe[
                "Source_File"
            ].iloc[0],

        "parsed_rows":
            len(dataframe),

        "unique_item_ids":
            int(
                dataframe[
                    "Item_ID"
                ]
                .nunique()
            ),

        "unique_descriptions":
            int(
                dataframe[
                    "Item_Description"
                ]
                .dropna()
                .nunique()
            ),

        "missing_description":
            int(
                dataframe[
                    "Item_Description"
                ]
                .isna()
                .sum()
            ),

        "missing_lpp":
            int(
                dataframe[
                    "Last_Purchase_Price"
                ]
                .isna()
                .sum()
            ),

        "missing_balance":
            int(
                dataframe[
                    "Balance_YE"
                ]
                .isna()
                .sum()
            ),

        "missing_total":
            int(
                dataframe[
                    "Total_Value_YE"
                ]
                .isna()
                .sum()
            ),

        "implied_cost_available":
            int(
                dataframe[
                    "Implied_Valuation_Cost"
                ]
                .notna()
                .sum()
            ),
    }


# ============================================================
# RECONCILIATION
# ============================================================

def reconcile_year(
    year_end: pd.DataFrame,
    annual: pd.DataFrame,
) -> tuple[
    pd.DataFrame,
    dict,
]:

    # Primary cross-source match:
    #
    # Year-end Item Description
    #       ->
    # Annual Brand_Name

    merged = year_end.merge(
        annual,
        left_on=(
            "Normalized_Description"
        ),
        right_on=(
            "Normalized_Brand"
        ),
        how="left",
        suffixes=(
            "_YE",
            "_Annual",
        ),
        indicator=True,
    )

    matched = (
        merged["_merge"]
        .eq("both")
    )

    # --------------------------------------------------------
    # GENERIC AGREEMENT
    # --------------------------------------------------------

    generic_comparable = (
        matched
        &
        merged[
            "Normalized_Generic_YE"
        ].notna()
        &
        merged[
            "Normalized_Generic_Annual"
        ].notna()
    )

    merged[
        "Generic_Match"
    ] = False

    merged.loc[
        generic_comparable,
        "Generic_Match",
    ] = (
        merged.loc[
            generic_comparable,
            "Normalized_Generic_YE",
        ]
        .eq(
            merged.loc[
                generic_comparable,
                "Normalized_Generic_Annual",
            ]
        )
    )

    # --------------------------------------------------------
    # UNIT AGREEMENT
    # --------------------------------------------------------

    unit_comparable = (
        matched
        &
        merged[
            "Normalized_Unit_YE"
        ].notna()
        &
        merged[
            "Normalized_Unit_Annual"
        ].notna()
    )

    merged[
        "Unit_Match"
    ] = False

    merged.loc[
        unit_comparable,
        "Unit_Match",
    ] = (
        merged.loc[
            unit_comparable,
            "Normalized_Unit_YE",
        ]
        .eq(
            merged.loc[
                unit_comparable,
                "Normalized_Unit_Annual",
            ]
        )
    )

    # --------------------------------------------------------
    # BALANCE AGREEMENT
    # --------------------------------------------------------

    balance_comparable = (
        matched
        &
        merged[
            "Balance_YE"
        ].notna()
        &
        merged[
            "Ending_Balance"
        ].notna()
    )

    merged[
        "Balance_Match"
    ] = False

    merged.loc[
        balance_comparable,
        "Balance_Match",
    ] = np.isclose(
        merged.loc[
            balance_comparable,
            "Balance_YE",
        ],
        merged.loc[
            balance_comparable,
            "Ending_Balance",
        ],
        atol=0.0001,
        rtol=0,
    )

    # --------------------------------------------------------
    # IMPLIED VALUATION COST
    # VS ANNUAL AVERAGE COST
    # --------------------------------------------------------

    implied_cost_comparable = (
        matched
        &
        merged[
            "Implied_Valuation_Cost"
        ].notna()
        &
        merged[
            "Average_cost"
        ].notna()
    )

    merged[
        "ImpliedCost_vs_AverageCost_Match"
    ] = False

    merged.loc[
        implied_cost_comparable,
        "ImpliedCost_vs_AverageCost_Match",
    ] = np.isclose(
        merged.loc[
            implied_cost_comparable,
            "Implied_Valuation_Cost",
        ],
        merged.loc[
            implied_cost_comparable,
            "Average_cost",
        ],
        atol=0.01,
        rtol=0,
    )

    # Absolute and percentage differences
    merged[
        "ImpliedCost_minus_AverageCost"
    ] = (
        merged[
            "Implied_Valuation_Cost"
        ]
        -
        merged[
            "Average_cost"
        ]
    )

    denominator = (
        merged[
            "Average_cost"
        ]
        .abs()
        .replace(
            0,
            np.nan,
        )
    )

    merged[
        "ImpliedCost_pct_difference"
    ] = (
        merged[
            "ImpliedCost_minus_AverageCost"
        ]
        .abs()
        /
        denominator
        * 100
    )

    # --------------------------------------------------------
    # LAST PURCHASE PRICE
    # VS ANNUAL AVERAGE COST
    # --------------------------------------------------------

    lpp_comparable = (
        matched
        &
        merged[
            "Last_Purchase_Price"
        ].notna()
        &
        merged[
            "Average_cost"
        ].notna()
    )

    merged[
        "LPP_vs_AverageCost_Match"
    ] = False

    merged.loc[
        lpp_comparable,
        "LPP_vs_AverageCost_Match",
    ] = np.isclose(
        merged.loc[
            lpp_comparable,
            "Last_Purchase_Price",
        ],
        merged.loc[
            lpp_comparable,
            "Average_cost",
        ],
        atol=0.01,
        rtol=0,
    )

    # --------------------------------------------------------
    # CAN YEAR-END IMPLIED COST POTENTIALLY
    # SUPPORT INVALID ANNUAL COST?
    # --------------------------------------------------------

    annual_invalid_cost = (
        merged[
            "Average_cost"
        ].isna()
        |
        merged[
            "Average_cost"
        ].le(0)
    )

    positive_issuance = (
        merged[
            "Qty_out_Stock_Issuances"
        ]
        .fillna(0)
        .gt(0)
    )

    merged[
        "Candidate_Cost_Support_Record"
    ] = (
        matched
        &
        annual_invalid_cost
        &
        positive_issuance
        &
        merged[
            "Implied_Valuation_Cost"
        ].notna()
        &
        merged[
            "Implied_Valuation_Cost"
        ].gt(0)
    )

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    matched_count = int(
        matched.sum()
    )

    summary = {
        "year":
            int(
                year_end[
                    "Year"
                ].iloc[0]
            ),

        "year_end_rows":
            len(year_end),

        "annual_medicine_rows":
            len(annual),

        "description_brand_matches":
            matched_count,

        "description_match_pct":
            percent(
                matched_count,
                len(year_end),
            ),

        "generic_comparable":
            int(
                generic_comparable.sum()
            ),

        "generic_matches":
            int(
                (
                    generic_comparable
                    &
                    merged[
                        "Generic_Match"
                    ]
                ).sum()
            ),

        "unit_comparable":
            int(
                unit_comparable.sum()
            ),

        "unit_matches":
            int(
                (
                    unit_comparable
                    &
                    merged[
                        "Unit_Match"
                    ]
                ).sum()
            ),

        "balance_comparable":
            int(
                balance_comparable.sum()
            ),

        "balance_matches":
            int(
                (
                    balance_comparable
                    &
                    merged[
                        "Balance_Match"
                    ]
                ).sum()
            ),

        "balance_match_pct":
            percent(
                int(
                    (
                        balance_comparable
                        &
                        merged[
                            "Balance_Match"
                        ]
                    ).sum()
                ),
                int(
                    balance_comparable.sum()
                ),
            ),

        "implied_cost_comparable":
            int(
                implied_cost_comparable.sum()
            ),

        "implied_cost_matches_average":
            int(
                (
                    implied_cost_comparable
                    &
                    merged[
                        "ImpliedCost_vs_AverageCost_Match"
                    ]
                ).sum()
            ),

        "implied_cost_match_pct":
            percent(
                int(
                    (
                        implied_cost_comparable
                        &
                        merged[
                            "ImpliedCost_vs_AverageCost_Match"
                        ]
                    ).sum()
                ),
                int(
                    implied_cost_comparable.sum()
                ),
            ),

        "lpp_comparable":
            int(
                lpp_comparable.sum()
            ),

        "lpp_matches_average":
            int(
                (
                    lpp_comparable
                    &
                    merged[
                        "LPP_vs_AverageCost_Match"
                    ]
                ).sum()
            ),

        "candidate_cost_support_records":
            int(
                merged[
                    "Candidate_Cost_Support_Record"
                ]
                .sum()
            ),
    }

    return (
        merged,
        summary,
    )


# ============================================================
# MAIN
# ============================================================

def main() -> None:

    print("=" * 92)
    print("CDH INVENTORY DSS")
    print("YEAR-END PHARMACY INVENTORY RECONCILIATION — CORRECTED")
    print("=" * 92)
    print()

    workbooks = sorted(
        YEAR_END_DIR.glob(
            "*.xlsx"
        )
    )

    all_parsed = []
    profiles = []
    reconciliation_frames = []
    summaries = []

    for workbook_path in workbooks:

        year = infer_year(
            workbook_path.name
        )

        print("-" * 92)
        print(
            f"YEAR: {year}"
        )

        print(
            f"FILE: {workbook_path.name}"
        )

        year_end = (
            parse_year_end_workbook(
                workbook_path
            )
        )

        annual = (
            load_annual_medicines(
                year
            )
        )

        profile = (
            profile_year_end(
                year_end
            )
        )

        profiles.append(
            profile
        )

        reconciliation, summary = (
            reconcile_year(
                year_end,
                annual,
            )
        )

        all_parsed.append(
            year_end
        )

        reconciliation_frames.append(
            reconciliation
        )

        summaries.append(
            summary
        )

        print()
        print(
            "Parsed MED rows:",
            f"{len(year_end):,}",
        )

        print(
            "Unique Item IDs:",
            f"{year_end['Item_ID'].nunique():,}",
        )

        print(
            "Unique descriptions:",
            f"{year_end['Item_Description'].nunique():,}",
        )

        print(
            "Implied valuation cost available:",
            f"{year_end['Implied_Valuation_Cost'].notna().sum():,}",
        )

        print()
        print(
            "First 5 parsed records:"
        )

        print(
            year_end[
                [
                    "Item_ID",
                    "Generic_Name_YE",
                    "Item_Description",
                    "Unit_YE",
                    "Last_Purchase_Price",
                    "Balance_YE",
                    "Total_Value_YE",
                    "Implied_Valuation_Cost",
                ]
            ]
            .head(5)
            .to_string(
                index=False
            )
        )

        print()
        print(
            "Description -> annual Brand_Name:"
        )

        print(
            "  Matched:",
            f"{summary['description_brand_matches']:,}",
            "/",
            f"{summary['year_end_rows']:,}",
            f"({summary['description_match_pct']}%)",
        )

        print()
        print(
            "Generic agreement:"
        )

        print(
            "  Comparable:",
            f"{summary['generic_comparable']:,}",
        )

        print(
            "  Matches:",
            f"{summary['generic_matches']:,}",
        )

        print()
        print(
            "Unit agreement:"
        )

        print(
            "  Comparable:",
            f"{summary['unit_comparable']:,}",
        )

        print(
            "  Matches:",
            f"{summary['unit_matches']:,}",
        )

        print()
        print(
            "Balance agreement:"
        )

        print(
            "  Comparable:",
            f"{summary['balance_comparable']:,}",
        )

        print(
            "  Matches annual Ending_Balance:",
            f"{summary['balance_matches']:,}",
            f"({summary['balance_match_pct']}%)",
        )

        print()
        print(
            "Derived implied valuation cost "
            "vs annual Average_cost:"
        )

        print(
            "  Comparable:",
            f"{summary['implied_cost_comparable']:,}",
        )

        print(
            "  Matches within ₱0.01:",
            f"{summary['implied_cost_matches_average']:,}",
            f"({summary['implied_cost_match_pct']}%)",
        )

        print()
        print(
            "Last Purchase Price "
            "vs annual Average_cost:"
        )

        print(
            "  Comparable:",
            f"{summary['lpp_comparable']:,}",
        )

        print(
            "  Matches within ₱0.01:",
            f"{summary['lpp_matches_average']:,}",
        )

        print()
        print(
            "Positive-issuance records with "
            "invalid annual cost but usable "
            "year-end implied valuation cost:"
        )

        print(
            " ",
            f"{summary['candidate_cost_support_records']:,}",
        )

        print()

    # ========================================================
    # SAVE
    # ========================================================

    parsed_df = pd.concat(
        all_parsed,
        ignore_index=True,
    )

    profile_df = pd.DataFrame(
        profiles
    )

    reconciliation_df = pd.concat(
        reconciliation_frames,
        ignore_index=True,
    )

    summary_df = pd.DataFrame(
        summaries
    )

    parsed_df.to_csv(
        OUTPUT_DIR
        / "year_end_pharmacy_parsed.csv",
        index=False,
        encoding="utf-8-sig",
    )

    profile_df.to_csv(
        OUTPUT_DIR
        / "year_end_pharmacy_profile.csv",
        index=False,
        encoding="utf-8-sig",
    )

    reconciliation_df.to_csv(
        OUTPUT_DIR
        / "year_end_annual_reconciliation_detail.csv",
        index=False,
        encoding="utf-8-sig",
    )

    summary_df.to_csv(
        OUTPUT_DIR
        / "year_end_annual_reconciliation_summary.csv",
        index=False,
        encoding="utf-8-sig",
    )

    print("=" * 92)
    print(
        "YEAR-END RECONCILIATION COMPLETE"
    )
    print("=" * 92)

    print()
    print(
        summary_df
        .to_string(
            index=False
        )
    )

    print()
    print(
        "Important:"
    )

    print(
        "  Last Purchase Price has NOT "
        "been renamed or treated as "
        "Average Cost."
    )

    print(
        "  Implied Valuation Cost is "
        "a DERIVED supporting value:"
    )

    print(
        "      Total_Value_YE / Balance_YE"
    )

    print(
        "  No annual cost has yet been "
        "replaced."
    )

    print(
        "  No raw source file was modified."
    )


if __name__ == "__main__":
    main()