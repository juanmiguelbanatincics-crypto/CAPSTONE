from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "year_end_inventory"
    / "dec 2022.xlsx"
)


def main():
    print("=" * 90)
    print("CDH INVENTORY DSS")
    print("YEAR-END EXCEL PHYSICAL LAYOUT INSPECTION")
    print("=" * 90)
    print()

    workbook = load_workbook(
        FILE_PATH,
        read_only=False,
        data_only=True,
    )

    worksheet = workbook[
        workbook.sheetnames[0]
    ]

    print("File:", FILE_PATH.name)
    print("Sheet:", worksheet.title)
    print(
        "Dimensions:",
        worksheet.max_row,
        "rows x",
        worksheet.max_column,
        "columns",
    )

    print()
    print("-" * 90)
    print("MERGED CELL RANGES")
    print("-" * 90)

    for merged_range in (
        worksheet.merged_cells.ranges
    ):
        print(merged_range)

    print()
    print("-" * 90)
    print("PHYSICAL ROWS 20 TO 45")
    print("-" * 90)

    for row_number in range(
        20,
        46,
    ):
        values = []

        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            value = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            if value is None:
                continue

            if isinstance(value, str):
                value = value.strip()

                if not value:
                    continue

            column_letter = (
                get_column_letter(
                    column_number
                )
            )

            text = str(value)

            if len(text) > 120:
                text = text[:117] + "..."

            values.append(
                f"{column_letter}{row_number}={text}"
            )

        print()
        print(
            f"ROW {row_number}:"
        )

        if values:
            for value in values:
                print(
                    " ",
                    value,
                )
        else:
            print(
                "  <EMPTY>"
            )

    print()
    print("-" * 90)
    print(
        "FIRST 10 ROWS AFTER HEADER "
        "WITH AT LEAST ONE VALUE"
    )
    print("-" * 90)

    displayed = 0

    for row_number in range(
        28,
        worksheet.max_row + 1,
    ):
        values = []

        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            value = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            if value is None:
                continue

            if isinstance(value, str):
                value = value.strip()

                if not value:
                    continue

            column_letter = (
                get_column_letter(
                    column_number
                )
            )

            text = str(value)

            if len(text) > 120:
                text = text[:117] + "..."

            values.append(
                f"{column_letter}{row_number}={text}"
            )

        if values:
            print()
            print(
                f"ROW {row_number}:"
            )

            for value in values:
                print(
                    " ",
                    value,
                )

            displayed += 1

        if displayed >= 10:
            break

    print()
    print("=" * 90)
    print(
        "LAYOUT INSPECTION COMPLETE"
    )
    print("=" * 90)

    print(
        "No workbook value was modified."
    )


if __name__ == "__main__":
    main()